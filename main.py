import os
import io
import re
import json
import requests
import openpyxl
from datetime import datetime
from fastapi import FastAPI, Form, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from openpyxl.styles import Font, PatternFill
# pyrefly: ignore [missing-import]
from pymongo import MongoClient, ReturnDocument

# -------------------------------
# MONGO & COUNTER STORAGE
# -------------------------------
MONGODB_URI = os.getenv(
    "MONGODB_URI",
    "mongodb+srv://sumaiyakn28_db_user:LxzSa6R77BOolLjt@cluster0.90gky9n.mongodb.net/drive_to_sheet?retryWrites=true&w=majority&appName=Cluster0"
)
COUNTER_FILE = os.path.join(os.path.dirname(__file__), "counter.json")

_mongo_client = None


def get_mongo_client():
    global _mongo_client
    if _mongo_client is None:
        try:
            _mongo_client = MongoClient(MONGODB_URI, serverSelectionTimeoutMS=4000)
        except Exception as e:
            print(f"[Counter] Failed to initialize MongoClient: {e}")
            _mongo_client = None
    return _mongo_client


def get_mongo_collection():
    client = get_mongo_client()
    if client is not None:
        try:
            client.admin.command('ping')
            db = client["drive_to_sheet"]
            return db["stats"]
        except Exception as e:
            print(f"[Counter] MongoDB fallback to local file: {e}")
    return None


def read_local_counter():
    if os.path.exists(COUNTER_FILE):
        try:
            with open(COUNTER_FILE, "r") as f:
                return json.load(f)
        except Exception:
            pass
    return {"total_generations": 0, "total_files_processed": 0, "last_generated_at": None}


def write_local_counter(data):
    try:
        with open(COUNTER_FILE, "w") as f:
            json.dump(data, f, indent=2)
    except Exception as e:
        print(f"[Counter] Failed to write local counter file: {e}")


def get_generation_stats():
    coll = get_mongo_collection()
    if coll is not None:
        try:
            doc = coll.find_one({"_id": "excel_counter"})
            if not doc:
                doc = {"_id": "excel_counter", "total_generations": 0, "total_files_processed": 0, "last_generated_at": None}
                coll.insert_one(doc)
            return {
                "total_generations": doc.get("total_generations", 0),
                "total_files_processed": doc.get("total_files_processed", 0),
                "last_generated_at": doc.get("last_generated_at")
            }
        except Exception as e:
            print(f"[Counter] Error fetching from MongoDB: {e}")
    return read_local_counter()


def increment_generation_stats(files_count: int = 0):
    now_str = datetime.utcnow().isoformat()
    coll = get_mongo_collection()
    new_stats = None
    if coll is not None:
        try:
            res = coll.find_one_and_update(
                {"_id": "excel_counter"},
                {
                    "$inc": {"total_generations": 1, "total_files_processed": files_count},
                    "$set": {"last_generated_at": now_str}
                },
                upsert=True,
                return_document=ReturnDocument.AFTER
            )
            if res:
                new_stats = {
                    "total_generations": res.get("total_generations", 0),
                    "total_files_processed": res.get("total_files_processed", 0),
                    "last_generated_at": res.get("last_generated_at")
                }
        except Exception as e:
            print(f"[Counter] Error updating MongoDB counter: {e}")
    
    local_data = read_local_counter()
    if new_stats:
        local_data.update(new_stats)
    else:
        local_data["total_generations"] = local_data.get("total_generations", 0) + 1
        local_data["total_files_processed"] = local_data.get("total_files_processed", 0) + files_count
        local_data["last_generated_at"] = now_str
    
    write_local_counter(local_data)
    return local_data


# -------------------------------
# FASTAPI APP
# -------------------------------
app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition", "X-Total-Generations"],
)


# -------------------------------
# EXTRACT ID FROM URL
# -------------------------------
def extract_id(url_or_id: str):
    url_or_id = url_or_id.strip().split("?")[0]

    match = re.search(r"/folders/([a-zA-Z0-9_-]{15,})", url_or_id)
    if match:
        return match.group(1)

    match_d = re.search(r"/d/([a-zA-Z0-9_-]{15,})", url_or_id)
    if match_d:
        return match_d.group(1)

    if re.match(r"^[a-zA-Z0-9_-]{15,}$", url_or_id):
        return url_or_id

    return url_or_id


# -------------------------------
# HOME & STATS
# -------------------------------
@app.get("/")
def home():
    return {
        "ok": True,
        "message": "Backend running on Render! API Key NOT required. Web scraping mode active."
    }


@app.get("/stats")
def get_stats():
    return get_generation_stats()


# -------------------------------
# FETCH PUBLIC FOLDER FILES
# -------------------------------
def get_public_folder_files(folder_id: str) -> list:
    url = f"https://drive.google.com/drive/folders/{folder_id}"
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/120.0.0.0 Safari/537.36"
        )
    }

    resp = requests.get(url, headers=headers, timeout=30)
    if resp.status_code != 200:
        raise Exception(f"Failed to load Drive folder page (HTTP {resp.status_code}). Make sure the folder is 'Anyone with the link can view'.")

    html = resp.text
    files = []
    seen_ids = set()

    def get_all_strings(o):
        res = []
        if isinstance(o, str):
            res.append(o)
        elif isinstance(o, list):
            for c in o:
                res.extend(get_all_strings(c))
        return res

    def process_item_node(item):
        if not isinstance(item, list) or len(item) == 0:
            return

        file_id = None
        # Primary location for Drive file item ID: item[0][1] or item[0][0][1]
        try:
            if isinstance(item[0], list) and len(item[0]) > 0:
                if len(item[0]) > 1 and isinstance(item[0][1], str) and re.match(r'^[a-zA-Z0-9_-]{25,50}$', item[0][1]):
                    file_id = item[0][1]
                elif isinstance(item[0][0], list) and len(item[0][0]) > 1 and isinstance(item[0][0][1], str) and re.match(r'^[a-zA-Z0-9_-]{25,50}$', item[0][0][1]):
                    file_id = item[0][0][1]
        except Exception:
            pass

        if not file_id:
            return

        if file_id == folder_id or file_id in seen_ids:
            return

        file_name = None
        # Primary structural location for item name: item[35][0][0][0]
        try:
            if isinstance(item[35], list) and len(item[35]) > 0:
                if isinstance(item[35][0], list) and len(item[35][0]) > 0:
                    if isinstance(item[35][0][0], list) and len(item[35][0][0]) > 0:
                        cand = item[35][0][0][0]
                        if isinstance(cand, str) and len(cand) > 0:
                            file_name = cand
        except Exception:
            pass

        # Secondary search across all strings in item
        if not file_name:
            strs = get_all_strings(item)
            candidates = [
                s for s in strs 
                if s != file_id 
                and not s.startswith("http") 
                and not s.startswith("image/") 
                and not s.startswith("application/") 
                and not s.startswith("video/") 
                and s not in ["Image", "Shared", "Folder", "Modified", "Download", "More actions", "1", ""]
                and not s.endswith(" KB") and not s.endswith(" MB") and not s.startswith("Size:")
                and len(s) > 1
            ]
            if candidates:
                file_name = candidates[0]

        if not file_name:
            file_name = f"File_{file_id}"

        files.append({"name": file_name, "id": file_id})
        seen_ids.add(file_id)

    # 1. Parse AF_initDataCallback JS payloads
    callbacks = re.findall(r'AF_initDataCallback\s*\(\s*({.*?})\s*\)\s*;', html, re.DOTALL)
    for cb in callbacks:
        match = re.search(r'data:\s*(\[.*\])\s*,\s*sideChannel:', cb, re.DOTALL)
        if not match:
            match = re.search(r'data:\s*(\[.*\])\s*\}\s*$', cb, re.DOTALL)
        if match:
            try:
                data = json.loads(match.group(1))
                def walk(obj):
                    if isinstance(obj, list):
                        process_item_node(obj)
                        for child in obj:
                            walk(child)
                    elif isinstance(obj, dict):
                        for k, v in obj.items():
                            walk(v)
                walk(data)
            except Exception:
                pass

    # 2. Fallback regex search if AF_initDataCallback missed files
    if not files:
        matches = re.findall(r'\[\[null,"([a-zA-Z0-9_-]{25,50})"\][^\]]*?,"([^"\\]+)"', html)
        for fid, fname in matches:
            if fid not in seen_ids and fid != folder_id and len(fname) < 200:
                files.append({"name": fname, "id": fid})
                seen_ids.add(fid)

    return files


# -------------------------------
# GENERATE EXCEL FROM GOOGLE DRIVE
# -------------------------------
@app.post("/generate-excel")
def generate_excel_from_drive(folder_id: str = Form(...)):
    try:
        folder_id_clean = extract_id(folder_id)
        files = get_public_folder_files(folder_id_clean)

        if not files:
            raise HTTPException(status_code=404, detail="No files found. Ensure the folder is not empty and shared as 'Anyone with the link can view'.")

        # Create Excel workbook in memory
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Drive Files"

        # Headers
        ws.append(["File Name", "File Link"])
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1155CC")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        for f in files:
            name = f["name"]
            link = f"https://drive.google.com/thumbnail?id={f['id']}"
            ws.append([name, link])

        ws.column_dimensions["A"].width = 60
        ws.column_dimensions["B"].width = 80

        # Save to BytesIO
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        # Increment generation stats in MongoDB & local counter
        stats = increment_generation_stats(len(files))

        # Return as StreamingResponse
        headers = {
            'Content-Disposition': 'attachment; filename="drive_files.xlsx"',
            'X-Total-Generations': str(stats.get("total_generations", 0))
        }
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers=headers
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Unexpected error: {str(e)}")


# -------------------------------
# RENDER START ENTRYPOINT
# -------------------------------
if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=int(os.getenv("PORT", 10000)))

